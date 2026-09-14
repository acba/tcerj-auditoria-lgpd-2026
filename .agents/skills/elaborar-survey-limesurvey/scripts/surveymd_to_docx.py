#!/usr/bin/env python3
"""Gera um guia DOCX legível para os respondentes a partir de SurveyMD."""

from __future__ import annotations

import argparse
import html
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from html.parser import HTMLParser
from pathlib import Path
from typing import Any, Iterable

SCRIPT_DIR = Path(__file__).resolve().parent
SKILL_DIR = SCRIPT_DIR.parent
DEFAULT_REFERENCE_DOCX = SKILL_DIR / "assets" / "template-questionario.docx"

if str(SCRIPT_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPT_DIR))

import surveymd_to_lss as engine  # noqa: E402


class _HTMLToText(HTMLParser):
    """Converte o HTML simples dos textos do survey em parágrafos legíveis."""

    BLOCKS = {"p", "div", "h1", "h2", "h3", "h4", "h5", "h6", "ul", "ol", "table", "tr"}

    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []
        self.skipped = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        tag = tag.lower()
        if tag in {"script", "style"}:
            self.skipped += 1
            return
        if self.skipped:
            return
        if tag == "br":
            self.parts.append("\n")
        elif tag == "li":
            self.parts.append("\n- ")
        elif tag == "img":
            alt = dict(attrs).get("alt")
            if alt:
                self.parts.append(f"[{alt}]")
        elif tag in self.BLOCKS:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        tag = tag.lower()
        if tag in {"script", "style"}:
            self.skipped = max(0, self.skipped - 1)
            return
        if not self.skipped and tag in self.BLOCKS | {"li"}:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if not self.skipped:
            self.parts.append(data)

    def text(self) -> str:
        value = html.unescape("".join(self.parts)).replace("\xa0", " ")
        value = re.sub(r"[ \t]+", " ", value)
        value = re.sub(r" *\n *", "\n", value)
        value = re.sub(r"\n{3,}", "\n\n", value)
        return value.strip()


def plain_text(value: object) -> str:
    text = str(value or "").strip()
    if "<" in text and ">" in text:
        parser = _HTMLToText()
        parser.feed(text)
        text = parser.text()
    text = re.sub(r"!\[([^]]*)\]\([^)]*\)", r"\1", text)
    text = re.sub(r"\[([^]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    text = text.replace("**", "").replace("__", "").replace("`", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    return text.strip()


def _all_questions(survey: engine.Survey) -> list[engine.Question]:
    return [question for group in survey.groups for question in group.questions]


def _question_index(survey: engine.Survey) -> dict[str, engine.Question]:
    return {question.code: question for question in _all_questions(survey)}


def _question_options(survey: engine.Survey, question: engine.Question) -> list[engine.Option]:
    if question.alternatives:
        return list(question.alternatives)
    if question.scale and question.scale in survey.scales:
        return list(survey.scales[question.scale].options)
    return []


def _option_text(survey: engine.Survey, question_code: str, option_code: str) -> str:
    question = _question_index(survey).get(question_code)
    if question is None:
        return option_code
    options = _question_options(survey, question) + list(question.subquestions)
    for option in options:
        if option.code == option_code:
            return plain_text(option.text)
    return option_code


def _simple_condition(survey: engine.Survey, clause: str) -> str | None:
    clause = clause.strip().strip("() ")
    checkbox = re.fullmatch(
        r"([A-Za-z][A-Za-z0-9_]*)[.\[]([A-Za-z0-9_]+)\]?\s*==\s*[\"']?Y[\"']?",
        clause,
        flags=re.I,
    )
    if checkbox:
        question_code, option_code = checkbox.groups()
        label = _option_text(survey, question_code, option_code)
        return f'o item “{label}” for marcado na pergunta {question_code}'

    membership = re.fullmatch(
        r"([A-Za-z][A-Za-z0-9_]*)\s+(not\s+in|in)\s+\[([^]]+)\]",
        clause,
        flags=re.I,
    )
    if membership:
        question_code, operator, values = membership.groups()
        labels = [
            f'“{_option_text(survey, question_code, item.strip().strip(chr(34) + chr(39)))}”'
            for item in values.split(",")
            if item.strip()
        ]
        joined = ", ".join(labels[:-1]) + (" ou " if len(labels) > 1 else "") + labels[-1]
        if operator.lower().startswith("not"):
            return f'a resposta à pergunta {question_code} não for {joined}'
        return f'a resposta à pergunta {question_code} for {joined}'

    comparison = re.fullmatch(
        r"([A-Za-z][A-Za-z0-9_]*)\s*(==|!=)\s*[\"']?([A-Za-z0-9_]+)[\"']?",
        clause,
        flags=re.I,
    )
    if comparison:
        question_code, operator, option_code = comparison.groups()
        label = _option_text(survey, question_code, option_code)
        verb = "for" if operator == "==" else "não for"
        return f'a resposta à pergunta {question_code} {verb} “{label}”'
    return None


def describe_visibility(survey: engine.Survey, expression: str, *, show_logic: bool = False) -> str:
    expression = (expression or "1").strip()
    if expression in {"", "1", "true", "True"}:
        return ""
    if expression.lower().startswith("raw:"):
        description = "Esta pergunta é exibida conforme respostas anteriores no formulário eletrônico."
    else:
        pieces = re.split(r"\s+(or|and)\s+", expression, flags=re.I)
        rendered: list[str] = []
        success = True
        for piece in pieces:
            lowered = piece.lower()
            if lowered == "or":
                rendered.append("ou")
            elif lowered == "and":
                rendered.append("e")
            else:
                item = _simple_condition(survey, piece)
                if item is None:
                    success = False
                    break
                rendered.append(item)
        if success:
            description = "Esta pergunta será exibida quando " + " ".join(rendered) + "."
        else:
            description = "Esta pergunta é exibida conforme respostas anteriores no formulário eletrônico."
    if show_logic:
        description += f" Condição técnica: {expression}"
    return description


def _escape_markdown(text: str) -> str:
    return text.replace("\\", "\\\\").replace("\n", "  \n")


def _metadata_lines(survey: engine.Survey, question: engine.Question, *, show_logic: bool) -> list[str]:
    # O reference DOCX institucional usa texto justificado e estilos de lista
    # compactos. Parágrafos separados evitam que metadados extensos aparentem
    # formar uma única linha no guia destinado aos respondentes.
    lines = [f"**Obrigatória:** {'Sim' if question.mandatory else 'Não'}", ""]
    visibility = describe_visibility(survey, question.visible_if, show_logic=show_logic)
    if visibility:
        lines.extend([f"*{visibility}*", ""])
    help_text = plain_text(question.help)
    if help_text:
        lines.extend([f"**Orientação:** {_escape_markdown(help_text)}", ""])
    return lines


def _append_choices(lines: list[str], choices: Iterable[str], *, marker: str = "•") -> None:
    """Acrescenta opções como parágrafos, sem depender do estilo List Bullet."""

    for choice in choices:
        lines.extend([f"{marker} {plain_text(choice)}", ""])


def render_respondent_markdown(
    survey: engine.Survey,
    *,
    source_name: str,
    show_logic: bool = False,
) -> str:
    title = plain_text(survey.meta.get("title") or "Questionário")
    admin = plain_text(survey.meta.get("admin") or "Equipe responsável")
    welcome = plain_text(survey.meta.get("welcome"))
    endtext = plain_text(survey.meta.get("endtext"))
    lines = [
        "---",
        f'title: "{title.replace(chr(34), chr(39))}"',
        'subtitle: "Documento auxiliar para preenchimento do formulário eletrônico"',
        f'author: "{admin.replace(chr(34), chr(39))}"',
        'lang: "pt-BR"',
        "---",
        "",
        "**ATENÇÃO:** este documento é um guia de consulta. O preenchimento e o envio das respostas devem ser realizados no LimeSurvey pelo link encaminhado à organização. Em caso de divergência, prevalecem os campos e as regras apresentados no formulário eletrônico.",
        "",
        f"*Fonte do questionário: {source_name}.*",
        "",
    ]
    if welcome:
        lines.extend(["# Orientações gerais", "", _escape_markdown(welcome), ""])

    for group in survey.groups:
        lines.extend([f"# {plain_text(group.title)}", ""])
        description = "\n".join(group.description_lines).strip()
        if description:
            lines.extend([description, ""])
        for question in group.questions:
            question_text = plain_text(question.text()) or question.code
            lines.extend([f"## {question_text}", "", f"*Código no formulário: {question.code}*", ""])
            lines.extend(_metadata_lines(survey, question, show_logic=show_logic))

            qtype = question.type.lower()
            options = _question_options(survey, question)
            if qtype in {"array", "matrix", "matriz", "array_numbers", "array_number", "numeric_array", "array_numeros", "matriz_numerica", "f", ":"}:
                if question.subquestions:
                    lines.extend(["**Itens a responder:**", ""])
                    _append_choices(lines, (item.text for item in question.subquestions))
                if options:
                    lines.extend(["**Alternativas disponíveis para cada item:**", ""])
                    _append_choices(lines, (item.text for item in options), marker="○")
            elif qtype in {"multi", "multiple", "checkbox", "multipla", "múltipla", "m"}:
                lines.extend(["**Marque todas as alternativas aplicáveis:**", ""])
                _append_choices(lines, (item.text for item in question.subquestions), marker="☐")
            elif qtype in {"multi_text", "multitext", "varios_textos", "q"}:
                lines.extend(["**Campos a preencher:**", ""])
                for item in question.subquestions:
                    lines.extend(
                        [
                            f"**{plain_text(item.text)}:**",
                            "",
                            "________________________________________",
                            "",
                        ]
                    )
            elif qtype in {"upload", "file", "arquivo", "|"}:
                formats = plain_text(question.attrs.get("allowed_filetypes", "arquivo permitido no formulário"))
                minimum = plain_text(question.attrs.get("min_files", ""))
                maximum = plain_text(question.attrs.get("max_files", ""))
                lines.extend([f"**Documento a anexar no LimeSurvey. Formatos:** {formats}.", ""])
                if minimum or maximum:
                    limits = []
                    if minimum:
                        limits.append(f"mínimo {minimum}")
                    if maximum:
                        limits.append(f"máximo {maximum}")
                    lines.extend([f"**Quantidade de arquivos:** {', '.join(limits)}.", ""])
            elif options:
                lines.extend(["**Selecione uma alternativa:**", ""])
                _append_choices(lines, (item.text for item in options), marker="○")
                if qtype in {"single_comment", "list_comment", "o"}:
                    lines.extend(["*Campo de comentário complementar disponível no formulário eletrônico.*", ""])
            elif qtype in {"long", "textarea", "texto_longo", "t"}:
                lines.extend(["**Campo de resposta:**", "", "________________________________________________________________________________", ""])
            else:
                lines.extend(["**Campo de resposta:** ________________________________________________", ""])

    if endtext:
        lines.extend(["# Encerramento", "", _escape_markdown(endtext), ""])
    return "\n".join(lines).rstrip() + "\n"


def _pandoc_command() -> tuple[str, Any] | None:
    try:
        import pypandoc

        pypandoc.get_pandoc_path()
        return "pypandoc", pypandoc
    except Exception:
        executable = shutil.which("pandoc")
        return ("cli", executable) if executable else None


def _convert_with_pandoc(
    markdown: str,
    output: Path,
    reference_docx: Path,
    resource_paths: Iterable[Path],
) -> str:
    command = _pandoc_command()
    if command is None:
        raise RuntimeError("Pandoc não está disponível.")
    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile("w", suffix=".md", encoding="utf-8", delete=False) as temp_file:
        temp_file.write(markdown)
        temp_path = Path(temp_file.name)
    try:
        args = [
            "--standalone",
            "--reference-doc=" + str(reference_docx),
            "--resource-path=" + os.pathsep.join(str(path) for path in resource_paths),
        ]
        kind, runner = command
        if kind == "pypandoc":
            runner.convert_file(str(temp_path), to="docx", outputfile=str(output), extra_args=args)
            return "pypandoc"
        subprocess.run([runner, str(temp_path), "-o", str(output), *args], check=True)
        return "pandoc"
    finally:
        temp_path.unlink(missing_ok=True)


def _set_run_font(run: Any, *, name: str, size: int, bold: bool | None = None, italic: bool | None = None, color: tuple[int, int, int] | None = None) -> None:
    from docx.shared import Pt, RGBColor

    run.font.name = name
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)


def _add_footer_page_number(document: Any) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    for section in document.sections:
        paragraph = section.footer.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run("Página ")
        field_begin = OxmlElement("w:fldChar")
        field_begin.set(qn("w:fldCharType"), "begin")
        instruction = OxmlElement("w:instrText")
        instruction.set(qn("xml:space"), "preserve")
        instruction.text = " PAGE "
        field_end = OxmlElement("w:fldChar")
        field_end.set(qn("w:fldCharType"), "end")
        run._r.append(field_begin)
        run._r.append(instruction)
        run._r.append(field_end)


def _add_docx_hyperlink(paragraph: Any, label: str, url: str) -> None:
    """Acrescenta hyperlink externo a um parágrafo do python-docx."""

    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE

    relationship_id = paragraph.part.relate_to(
        url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.extend([color, underline])
    text = OxmlElement("w:t")
    text.text = label
    run.extend([properties, text])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _append_markdown_inline(paragraph: Any, value: str) -> None:
    """Renderiza negrito, itálico, código e links Markdown no python-docx."""

    token_re = re.compile(
        r"(\*\*[^*\n]+\*\*|__[^_\n]+__|`[^`\n]+`|"
        r"\[[^]\n]+\]\(https?://[^)\s]+\)|(?<!\*)\*[^*\n]+\*(?!\*))"
    )
    cursor = 0
    for match in token_re.finditer(value):
        if match.start() > cursor:
            paragraph.add_run(value[cursor:match.start()])
        token = match.group(0)
        link = re.fullmatch(r"\[([^]\n]+)\]\((https?://[^)\s]+)\)", token)
        if link:
            _add_docx_hyperlink(paragraph, link.group(1), link.group(2))
        elif token.startswith(("**", "__")):
            paragraph.add_run(token[2:-2]).bold = True
        elif token.startswith("*"):
            paragraph.add_run(token[1:-1]).italic = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Courier New"
        cursor = match.end()
    if cursor < len(value):
        paragraph.add_run(value[cursor:])


def _add_markdown_blocks(document: Any, markdown: str) -> None:
    """Acrescenta parágrafos e listas Markdown preservando formatação inline."""

    for block in re.split(r"\n\s*\n", markdown.strip()):
        lines = [line.strip() for line in block.splitlines() if line.strip()]
        if not lines:
            continue
        unordered = all(re.match(r"^[-+*]\s+.+$", line) for line in lines)
        ordered = all(re.match(r"^\d+[.)]\s+.+$", line) for line in lines)
        if unordered or ordered:
            pattern = r"^[-+*]\s+(.+)$" if unordered else r"^\d+[.)]\s+(.+)$"
            style = "List Bullet" if unordered else "List Number"
            for line in lines:
                match = re.match(pattern, line)
                paragraph = document.add_paragraph(style=style)
                _append_markdown_inline(paragraph, match.group(1) if match else line)
            continue
        paragraph = document.add_paragraph()
        for index, line in enumerate(lines):
            if index:
                paragraph.add_run().add_break()
            _append_markdown_inline(paragraph, line)


def _convert_with_python_docx(survey: engine.Survey, output: Path, *, source_name: str, show_logic: bool) -> str:
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Cm, Pt, RGBColor
    except ImportError as exc:
        raise RuntimeError("Instale pypandoc/Pandoc ou python-docx para gerar DOCX.") from exc

    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.0)
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    for style_name, size in (("Title", 20), ("Heading 1", 16), ("Heading 2", 12)):
        style = document.styles[style_name]
        style.font.name = "Cambria" if style_name != "Heading 2" else "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(47, 84, 150) if style_name != "Heading 2" else RGBColor(0, 0, 0)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run(plain_text(survey.meta.get("title") or "Questionário"))
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Documento auxiliar para preenchimento do formulário eletrônico")
    _set_run_font(run, name="Arial", size=12, bold=True)
    admin = plain_text(survey.meta.get("admin"))
    if admin:
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(admin)
    warning = document.add_paragraph()
    run = warning.add_run(
        "ATENÇÃO: este documento é um guia de consulta. O preenchimento e o envio das respostas devem ser realizados no LimeSurvey pelo link encaminhado à organização. Em caso de divergência, prevalecem os campos e as regras apresentados no formulário eletrônico."
    )
    _set_run_font(run, name="Arial", size=10, bold=True)
    source = document.add_paragraph()
    run = source.add_run(f"Fonte do questionário: {source_name}.")
    _set_run_font(run, name="Arial", size=9, italic=True)

    welcome = plain_text(survey.meta.get("welcome"))
    if welcome:
        document.add_heading("Orientações gerais", level=1)
        for paragraph in welcome.split("\n"):
            if paragraph.strip():
                document.add_paragraph(paragraph.strip())

    for group in survey.groups:
        document.add_heading(plain_text(group.title), level=1)
        description = "\n".join(group.description_lines).strip()
        if description:
            _add_markdown_blocks(document, description)
        for question in group.questions:
            document.add_heading(plain_text(question.text()) or question.code, level=2)
            code_p = document.add_paragraph()
            run = code_p.add_run(f"Código no formulário: {question.code}")
            _set_run_font(run, name="Arial", size=9, italic=True, color=(89, 89, 89))
            required = document.add_paragraph()
            required.add_run("Obrigatória: ").bold = True
            required.add_run("Sim" if question.mandatory else "Não")
            visibility = describe_visibility(survey, question.visible_if, show_logic=show_logic)
            if visibility:
                p = document.add_paragraph(visibility)
                for run in p.runs:
                    run.italic = True
            if question.help:
                p = document.add_paragraph()
                p.add_run("Orientação: ").bold = True
                p.add_run(plain_text(question.help))

            qtype = question.type.lower()
            options = _question_options(survey, question)
            if qtype in {"array", "matrix", "matriz", "array_numbers", "array_number", "numeric_array", "array_numeros", "matriz_numerica", "f", ":"}:
                if question.subquestions:
                    document.add_paragraph("Itens a responder:").runs[0].bold = True
                    for item in question.subquestions:
                        document.add_paragraph(plain_text(item.text), style="List Bullet")
                if options:
                    document.add_paragraph("Alternativas disponíveis para cada item:").runs[0].bold = True
                    for item in options:
                        document.add_paragraph(plain_text(item.text), style="List Bullet")
            elif qtype in {"multi", "multiple", "checkbox", "multipla", "múltipla", "m"}:
                document.add_paragraph("Marque todas as alternativas aplicáveis:").runs[0].bold = True
                for item in question.subquestions:
                    document.add_paragraph(f"☐ {plain_text(item.text)}", style="List Bullet")
            elif qtype in {"multi_text", "multitext", "varios_textos", "q"}:
                document.add_paragraph("Campos a preencher:").runs[0].bold = True
                for item in question.subquestions:
                    document.add_paragraph(f"{plain_text(item.text)}: ________________________________________", style="List Bullet")
            elif qtype in {"upload", "file", "arquivo", "|"}:
                p = document.add_paragraph()
                p.add_run("Documento a anexar no LimeSurvey. Formatos: ").bold = True
                p.add_run(plain_text(question.attrs.get("allowed_filetypes", "arquivo permitido no formulário")) + ".")
            elif options:
                document.add_paragraph("Selecione uma alternativa:").runs[0].bold = True
                for item in options:
                    document.add_paragraph(f"○ {plain_text(item.text)}", style="List Bullet")
                if qtype in {"single_comment", "list_comment", "o"}:
                    document.add_paragraph("Campo de comentário complementar disponível no formulário eletrônico.")
            elif qtype in {"long", "textarea", "texto_longo", "t"}:
                document.add_paragraph("Campo de resposta:\n\n________________________________________________________________________________")
            else:
                document.add_paragraph("Campo de resposta: ________________________________________________")

    endtext = plain_text(survey.meta.get("endtext"))
    if endtext:
        document.add_heading("Encerramento", level=1)
        document.add_paragraph(endtext)
    _add_footer_page_number(document)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    return "python-docx"


def validate_docx(path: Path, *, expected_text: str = "") -> dict[str, Any]:
    errors: list[str] = []
    if not path.exists() or path.stat().st_size == 0:
        errors.append("Arquivo DOCX ausente ou vazio.")
        return {"source": str(path), "valid": False, "errors": errors, "size": 0}
    try:
        with zipfile.ZipFile(path) as archive:
            names = set(archive.namelist())
            for required in {"[Content_Types].xml", "word/document.xml", "word/styles.xml"}:
                if required not in names:
                    errors.append(f"Parte obrigatória ausente no DOCX: {required}.")
            document_xml = archive.read("word/document.xml").decode("utf-8", errors="replace")
            if expected_text and expected_text not in document_xml:
                errors.append("O título esperado não foi encontrado no conteúdo do DOCX.")
            for directive in ("mandatory:", "visible_if:", "evidence_text:", "allowed_filetypes:"):
                if directive in document_xml:
                    errors.append(f"Diretiva técnica exposta no documento: {directive}")
    except (zipfile.BadZipFile, KeyError) as exc:
        errors.append(f"Pacote DOCX inválido: {exc}")
    return {
        "source": str(path),
        "valid": not errors,
        "errors": errors,
        "size": path.stat().st_size,
    }


def build_docx(
    source: Path,
    output: Path,
    *,
    reference_docx: Path = DEFAULT_REFERENCE_DOCX,
    resource_paths: Iterable[Path] = (),
    show_logic: bool = False,
    conversion_engine: str = "auto",
    target: str = "",
) -> dict[str, Any]:
    survey = engine.parse_markdown(source)
    if target:
        survey = engine.filter_survey_by_target(survey, target)
    markdown = render_respondent_markdown(survey, source_name=source.name, show_logic=show_logic)
    selected_engine = conversion_engine
    if selected_engine == "auto":
        selected_engine = "pandoc" if _pandoc_command() is not None else "python-docx"
    if selected_engine == "pandoc":
        if not reference_docx.exists():
            raise FileNotFoundError(f"Reference DOCX não encontrado: {reference_docx}")
        used_engine = _convert_with_pandoc(
            markdown,
            output,
            reference_docx,
            [source.parent, reference_docx.parent, *resource_paths],
        )
    elif selected_engine == "python-docx":
        used_engine = _convert_with_python_docx(
            survey,
            output,
            source_name=source.name,
            show_logic=show_logic,
        )
    else:
        raise ValueError("Engine inválida; use auto, pandoc ou python-docx.")
    validation = validate_docx(output, expected_text=plain_text(survey.meta.get("title") or "Questionário"))
    return {
        "source": str(source),
        "output": str(output),
        "engine": used_engine,
        "groups": len(survey.groups),
        "questions": len(_all_questions(survey)),
        "validation": validation,
    }


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--reference-docx", type=Path, default=DEFAULT_REFERENCE_DOCX)
    parser.add_argument("--resource-files", nargs="*", type=Path, default=[])
    parser.add_argument("--show-logic", action="store_true")
    parser.add_argument("--engine", choices=["auto", "pandoc", "python-docx"], default="auto")
    parser.add_argument("--target", default="")
    args = parser.parse_args(argv)
    try:
        result = build_docx(
            args.source,
            args.output,
            reference_docx=args.reference_docx,
            resource_paths=args.resource_files,
            show_logic=args.show_logic,
            conversion_engine=args.engine,
            target=args.target,
        )
    except Exception as exc:
        print(f"ERRO: {exc}", file=sys.stderr)
        return 1
    print(f"OK: DOCX gerado em {result['output']} ({result['engine']}).")
    print(f"Grupos: {result['groups']}; perguntas: {result['questions']}.")
    if not result["validation"]["valid"]:
        for error in result["validation"]["errors"]:
            print(f"ERRO: {error}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
