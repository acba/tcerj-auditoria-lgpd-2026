#!/usr/bin/env python3
"""Gera o guia institucional do questionário LGPD a partir do SurveyMD.

O documento de referência fornece página A4, estilos, cabeçalho, rodapé e a
identidade visual institucional. O conteúdo é sempre reconstruído a partir do
Markdown canônico do survey.
"""

from __future__ import annotations

import argparse
import os
import re
import sys
import tempfile
import zipfile
from datetime import datetime
from html.parser import HTMLParser
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[1]
SURVEY_SKILL_SCRIPTS = (
    PROJECT_ROOT / ".agents" / "skills" / "elaborar-survey-limesurvey" / "scripts"
)
if str(SURVEY_SKILL_SCRIPTS) not in sys.path:
    sys.path.insert(0, str(SURVEY_SKILL_SCRIPTS))

import surveymd_to_docx as survey_docx  # noqa: E402
import surveymd_to_lss as survey_engine  # noqa: E402


BLUE = "2F5597"
LIGHT_BLUE = "EAF2FB"
MID_BLUE = "5B9BD5"
PALE_BLUE = "F3F7FC"
RED = "C00000"
TEXT = "202020"
MUTED = "667085"
LINE = "D7E2F0"
FIXED_TIMESTAMP = (1980, 1, 1, 0, 0, 0)


class BlockHTMLParser(HTMLParser):
    """Extrai parágrafos e itens de lista na ordem em que aparecem."""

    def __init__(self) -> None:
        super().__init__()
        self.blocks: list[tuple[str, str]] = []
        self._kind = ""
        self._parts: list[str] = []
        self._skip = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        tag = tag.lower()
        if tag in {"script", "style"}:
            self._skip += 1
            return
        if self._skip:
            return
        if tag in {"p", "li"}:
            self._flush()
            self._kind = tag
        elif tag == "br" and self._kind:
            self._parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        tag = tag.lower()
        if tag in {"script", "style"}:
            self._skip = max(0, self._skip - 1)
            return
        if not self._skip and tag == self._kind:
            self._flush()

    def handle_data(self, data: str) -> None:
        if not self._skip and self._kind:
            self._parts.append(data)

    def close(self) -> None:
        super().close()
        self._flush()

    def _flush(self) -> None:
        if self._kind and self._parts:
            value = re.sub(r"\s+", " ", "".join(self._parts)).strip()
            if value:
                self.blocks.append((self._kind, value))
        self._kind = ""
        self._parts = []


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path, help="SurveyMD canônico.")
    parser.add_argument("output", type=Path, help="DOCX institucional de saída.")
    parser.add_argument(
        "--reference-docx",
        required=True,
        type=Path,
        help="DOCX institucional usado como referência de estilos, cabeçalho e rodapé.",
    )
    return parser.parse_args()


def clear_body(document: Document) -> None:
    body = document._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = tr_pr.find(qn("w:cantSplit"))
    if cant_split is None:
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, *, top: int = 100, start: int = 120, bottom: int = 100, end: int = 120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_twips: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    width = tc_pr.find(qn("w:tcW"))
    if width is None:
        width = OxmlElement("w:tcW")
        tc_pr.append(width)
    width.set(qn("w:w"), str(width_twips))
    width.set(qn("w:type"), "dxa")


def set_fixed_table_layout(table) -> None:
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")


def set_table_borders(table, *, color: str = LINE, size: str = "4") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def add_bottom_rule(paragraph, *, color: str = LINE, size: str = "6") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "5")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)


def set_run(run, *, size: float | None = None, bold: bool | None = None,
            italic: bool | None = None, color: str | None = None,
            name: str = "Arial") -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def iter_container_paragraphs(container):
    yield from container.paragraphs
    for table in container.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from iter_container_paragraphs(cell)


def configure_headers(document: Document) -> None:
    """Uniformiza em 8 pt o texto dos cabeçalhos herdados do modelo."""

    seen_parts: set[int] = set()
    for section in document.sections:
        for header in (section.header, section.even_page_header, section.first_page_header):
            part_id = id(header.part)
            if part_id in seen_parts:
                continue
            seen_parts.add(part_id)
            for paragraph in iter_container_paragraphs(header):
                for run in paragraph.runs:
                    set_run(run, size=8)


def configure_document(document: Document, *, title: str) -> None:
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(1.6)
    section.right_margin = Cm(1.6)
    section.top_margin = Cm(3.1)
    section.bottom_margin = Cm(2.1)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    title_style = document.styles["Title"]
    title_style.font.name = "Arial"
    title_style.font.size = Pt(21)
    title_style.font.bold = False
    title_style.font.color.rgb = RGBColor.from_string(BLUE)
    title_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_style.paragraph_format.space_after = Pt(0)

    for name, size, color in (
        ("Heading 1", 15, TEXT),
        ("Heading 2", 11, TEXT),
        ("Heading 3", 10.5, BLUE),
    ):
        style = document.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(8 if name == "Heading 1" else 6)
        style.paragraph_format.space_after = Pt(7 if name == "Heading 1" else 4)
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    custom_styles = {
        "Survey Metadata": (8.5, False, True, MUTED),
        "Survey Explain": (9.5, False, False, MID_BLUE),
        "Survey Help": (9, False, False, RED),
        "Survey Condition": (9, False, True, MUTED),
        "Survey Option": (9.7, False, False, TEXT),
        "Survey Label": (9.5, True, False, TEXT),
        "Survey Group Intro": (9.6, False, False, TEXT),
    }
    for name, (size, bold, italic, color) in custom_styles.items():
        if name in document.styles:
            style = document.styles[name]
        else:
            style = document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = normal
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.italic = italic
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.line_spacing = 1.03

    # O modelo institucional usa nomes localizados para listas. Criamos aliases
    # estáveis para que o gerador não dependa do idioma do Word/LibreOffice.
    for name in ("List Bullet", "List Number"):
        if name in document.styles:
            style = document.styles[name]
        else:
            style = document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = normal
        style.font.name = "Arial"
        style.font.size = Pt(10.5)
        style.paragraph_format.left_indent = Cm(0.65)
        style.paragraph_format.first_line_indent = Cm(-0.35)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.05

    explain_style = document.styles["Survey Explain"]
    explain_style.font.name = "Calibri"
    explain_style.font.size = Pt(9)
    explain_style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    explain_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    explain_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")

    document.core_properties.title = title
    document.core_properties.subject = "Documento auxiliar para preenchimento do LimeSurvey"
    document.core_properties.author = "Tribunal de Contas do Estado do Rio de Janeiro"
    document.core_properties.last_modified_by = "TCE-RJ"
    document.core_properties.created = datetime(2026, 1, 1)
    document.core_properties.modified = datetime(2026, 1, 1)
    configure_headers(document)


def html_blocks(value: object) -> list[tuple[str, str]]:
    parser = BlockHTMLParser()
    parser.feed(str(value or ""))
    parser.close()
    ignored_prefixes = (
        "SECRETARIA-GERAL DE CONTROLE EXTERNO",
        "SUBSECRETARIA DE CONTROLE",
        "COORDENADORIA DE AUDITORIA",
        "Segue documento para auxiliar",
        "Questionário em formato PDF",
    )
    return [
        (kind, text)
        for kind, text in parser.blocks
        if not text.startswith(ignored_prefixes)
    ]


def add_cover(document: Document, survey) -> None:
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_before = Pt(205)
    spacer.paragraph_format.space_after = Pt(14)

    annex = document.add_paragraph()
    annex.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = annex.add_run("ANEXO")
    set_run(run, size=17, color=BLUE)
    annex.paragraph_format.space_after = Pt(18)

    title = document.add_paragraph(style="Title")
    title.add_run(survey_docx.plain_text(survey.meta.get("title") or "Questionário"))

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_before = Pt(12)
    subtitle.paragraph_format.space_after = Pt(0)
    run = subtitle.add_run("Documento auxiliar para preenchimento do formulário eletrônico")
    set_run(run, size=10.5, color=MUTED)
    document.add_page_break()


def add_intro(document: Document, survey) -> None:
    document.add_heading("Introdução", level=1)
    for kind, text in html_blocks(survey.meta.get("welcome")):
        if text == "Observações importantes:":
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(7)
            paragraph.add_run(text).bold = True
            continue
        if kind == "li":
            paragraph = document.add_paragraph(style="List Number")
        else:
            paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        paragraph.add_run(text)

def add_group_description(document: Document, description_lines: list[str]) -> None:
    markdown = "\n".join(description_lines).strip()
    if not markdown:
        return
    cleaned = "\n".join(re.sub(r"^>\s?", "", line) for line in markdown.splitlines()).strip()
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_table_borders(table, color=LINE, size="4")
    cell = table.cell(0, 0)
    set_cell_shading(cell, PALE_BLUE)
    set_cell_margins(cell, top=120, bottom=110, start=150, end=150)
    first = cell.paragraphs[0]
    first.style = document.styles["Survey Group Intro"]
    first._element.getparent().remove(first._element)

    for block in re.split(r"\n\s*\n", cleaned):
        lines = [line.strip() for line in block.splitlines() if line.strip()]
        if not lines:
            continue
        unordered = all(re.match(r"^[-+*]\s+.+$", line) for line in lines)
        ordered = all(re.match(r"^\d+[.)]\s+.+$", line) for line in lines)
        if unordered or ordered:
            pattern = r"^[-+*]\s+(.+)$" if unordered else r"^(\d+)[.)]\s+(.+)$"
            for index, line in enumerate(lines, start=1):
                match = re.match(pattern, line)
                paragraph = cell.add_paragraph(style="Survey Group Intro")
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                paragraph.paragraph_format.left_indent = Cm(0.45)
                paragraph.paragraph_format.first_line_indent = Cm(-0.3)
                if unordered:
                    text = match.group(1) if match else line
                    paragraph.add_run("•  ")
                else:
                    number = match.group(1) if match else str(index)
                    text = match.group(2) if match else line
                    paragraph.add_run(f"{number}.  ")
                survey_docx._append_markdown_inline(paragraph, text)
            continue

        paragraph = cell.add_paragraph(style="Survey Group Intro")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for index, line in enumerate(lines):
            if index:
                paragraph.add_run().add_break()
            survey_docx._append_markdown_inline(paragraph, line)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_metadata(document: Document, question) -> None:
    if not question.mandatory:
        return
    paragraph = document.add_paragraph(style="Survey Metadata")
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run("resposta obrigatória")


def add_explain(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="Survey Explain")
    paragraph.paragraph_format.left_indent = Inches(1.5)
    paragraph.paragraph_format.right_indent = Cm(0.2)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(survey_docx.plain_text(text))
    set_run(run, size=9, name="Calibri", color=MID_BLUE)


def add_help(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="Survey Help")
    paragraph.paragraph_format.left_indent = Cm(0.1)
    paragraph.paragraph_format.right_indent = Cm(0.2)
    paragraph.paragraph_format.space_before = Pt(4)
    run = paragraph.add_run("?  ")
    set_run(run, size=9, bold=True, color=RED)
    paragraph.add_run(survey_docx.plain_text(text))


def add_option(document: Document, symbol: str, text: str) -> None:
    paragraph = document.add_paragraph(style="Survey Option")
    paragraph.paragraph_format.left_indent = Cm(0.65)
    paragraph.paragraph_format.first_line_indent = Cm(-0.42)
    paragraph.add_run(f"{symbol}  ")
    paragraph.add_run(survey_docx.plain_text(text))


def add_answer_lines(document: Document, *, count: int = 3) -> None:
    for _ in range(count):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(2)
        paragraph.paragraph_format.space_after = Pt(2)
        add_bottom_rule(paragraph, color="B8B8B8", size="3")


def add_upload_callout(document: Document, question) -> None:
    table = document.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_fixed_table_layout(table)
    table.columns[0].width = Cm(1.35)
    table.columns[1].width = Cm(16.25)
    set_table_borders(table, color=MID_BLUE, size="5")
    prevent_row_split(table.rows[0])
    icon_cell, text_cell = table.rows[0].cells
    set_cell_width(icon_cell, 765)
    set_cell_width(text_cell, 9212)
    set_cell_shading(icon_cell, MID_BLUE)
    set_cell_shading(text_cell, LIGHT_BLUE)
    set_cell_margins(icon_cell, top=90, bottom=90, start=70, end=70)
    set_cell_margins(text_cell, top=90, bottom=90, start=130, end=130)
    icon_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    text_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    icon = icon_cell.paragraphs[0]
    icon.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = icon.add_run("DOC")
    set_run(run, size=8, bold=True, color="FFFFFF")

    paragraph = text_cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run("Evidência documental")
    set_run(run, size=9.5, bold=True, color=BLUE)
    formats = survey_docx.plain_text(question.attrs.get("allowed_filetypes", "arquivo permitido"))
    detail = paragraph.add_run(f"\nEnvio pelo LimeSurvey · Formatos permitidos: {formats.upper()}.")
    set_run(detail, size=9, color=TEXT)


def add_question(document: Document, survey, question) -> None:
    heading = document.add_paragraph(style="Heading 2")
    heading.paragraph_format.keep_with_next = True
    heading.add_run(survey_docx.plain_text(question.text()) or question.code)
    add_metadata(document, question)

    visibility = survey_docx.describe_visibility(survey, question.visible_if, show_logic=False)
    if visibility:
        paragraph = document.add_paragraph(style="Survey Condition")
        paragraph.paragraph_format.keep_with_next = True
        paragraph.add_run(visibility)
    explain = question.attrs.get("explain", "").strip()
    if explain:
        add_explain(document, explain)

    qtype = question.type.lower()
    options = survey_docx._question_options(survey, question)
    if qtype in {"array", "matrix", "matriz", "array_numbers", "array_number", "numeric_array", "array_numeros", "matriz_numerica", "f", ":"}:
        if question.subquestions:
            document.add_paragraph("Itens a responder", style="Survey Label")
            for item in question.subquestions:
                add_option(document, "□", item.text)
        if options:
            document.add_paragraph("Alternativas disponíveis para cada item", style="Survey Label")
            for item in options:
                add_option(document, "○", item.text)
    elif qtype in {"multi", "multiple", "checkbox", "multipla", "múltipla", "m"}:
        for item in question.subquestions:
            add_option(document, "□", item.text)
    elif qtype in {"multi_text", "multitext", "varios_textos", "q"}:
        for item in question.subquestions:
            paragraph = document.add_paragraph(style="Survey Option")
            paragraph.add_run(survey_docx.plain_text(item.text) + ": ")
            add_bottom_rule(paragraph, color="B8B8B8", size="3")
    elif qtype in {"upload", "file", "arquivo", "|"}:
        add_upload_callout(document, question)
    elif options:
        for item in options:
            add_option(document, "○", item.text)
        if qtype in {"single_comment", "list_comment", "o"}:
            paragraph = document.add_paragraph(style="Survey Metadata")
            paragraph.add_run("Comentário complementar disponível no formulário eletrônico.")
    elif qtype in {"long", "textarea", "texto_longo", "t"}:
        add_answer_lines(document, count=3)
    else:
        add_answer_lines(document, count=1)

    if question.help:
        add_help(document, question.help)

    separator = document.add_paragraph()
    separator.paragraph_format.space_before = Pt(4)
    separator.paragraph_format.space_after = Pt(7)
    add_bottom_rule(separator, color=LINE, size="3")


def add_groups(document: Document, survey) -> None:
    for group in survey.groups:
        document.add_page_break()
        heading = document.add_paragraph(style="Heading 1")
        heading.add_run(survey_docx.plain_text(group.title))
        add_bottom_rule(heading, color=BLUE, size="8")
        add_group_description(document, group.description_lines)
        for question in group.questions:
            add_question(document, survey, question)


def add_end(document: Document, survey) -> None:
    document.add_page_break()
    document.add_heading("Encerramento", level=1)
    for kind, text in html_blocks(survey.meta.get("endtext")):
        paragraph = document.add_paragraph(style="List Number" if kind == "li" else None)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        paragraph.add_run(text)


def normalize_docx(path: Path) -> None:
    """Normaliza metadados ZIP para tornar a saída byte a byte reproduzível."""

    with zipfile.ZipFile(path, "r") as source:
        entries = [(item.filename, source.read(item.filename)) for item in source.infolist()]
    descriptor, temporary_name = tempfile.mkstemp(suffix=".docx", dir=path.parent)
    os.close(descriptor)
    temporary = Path(temporary_name)
    try:
        with zipfile.ZipFile(temporary, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as target:
            for name, data in sorted(entries):
                info = zipfile.ZipInfo(name, FIXED_TIMESTAMP)
                info.compress_type = zipfile.ZIP_DEFLATED
                info.external_attr = 0o600 << 16
                target.writestr(info, data)
        os.replace(temporary, path)
    finally:
        if temporary.exists():
            temporary.unlink()


def main() -> int:
    args = parse_args()
    if not args.source.is_file():
        raise SystemExit(f"SurveyMD não encontrado: {args.source}")
    if not args.reference_docx.is_file():
        raise SystemExit(f"DOCX de referência não encontrado: {args.reference_docx}")

    survey = survey_engine.parse_markdown(args.source)
    document = Document(args.reference_docx)
    clear_body(document)
    title = survey_docx.plain_text(survey.meta.get("title") or "Questionário")
    configure_document(document, title=title)
    add_cover(document, survey)
    add_intro(document, survey)
    add_groups(document, survey)
    add_end(document, survey)

    args.output.parent.mkdir(parents=True, exist_ok=True)
    document.save(args.output)
    normalize_docx(args.output)

    validation = survey_docx.validate_docx(args.output, expected_text=title)
    if not validation["valid"]:
        for error in validation["errors"]:
            print(f"ERRO: {error}", file=sys.stderr)
        return 1
    print(f"DOCX: {args.output}")
    print(f"Grupos: {len(survey.groups)} | Perguntas: {sum(len(g.questions) for g in survey.groups)}")
    print(f"Tamanho: {validation['size']} bytes")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
